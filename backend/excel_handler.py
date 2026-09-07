from openpyxl import Workbook, load_workbook
import os
from backend.text_extractor import extract_text
from backend.llm_parser import parse_json

# ========== File Path & Extracted Text & JSON Data ==========
file_path = r"D:\HIMANSHU\Desktop\RD\Project-8 Invoice Data Parser\backend\data\batch1_1\batch1-0001.jpg"
text = extract_text(file_path)
data = parse_json(text)


INVOICE_HEADERS = [
    "Invoice No",
    "Date",
    "Seller",
    "Seller Address",
    "Seller Tax ID",
    "Client",
    "Client Address",
    "Client Tax ID",
    "Net Total",
    "VAT",
    "Gross Total"
]

ITEM_HEADERS = [
    "Invoice No",
    "Item No",
    "Description",
    "Quantity",
    "Unit",
    "Net Price",
    "Net Worth",
    "VAT %",
    "Gross Worth"
]


def save_to_excel(data, file_path="invoices.xlsx"):

    # ========== Open Existing Excel or Creating New One ==========
    if os.path.exists(file_path):

        workbook = load_workbook(file_path)

        if "Invoices" in workbook.sheetnames:
            invoice_sheet = workbook["Invoices"]
        else:
            invoice_sheet = workbook.create_sheet("Invoices")
            invoice_sheet.append(INVOICE_HEADERS)

        if "Items" in workbook.sheetnames:
            item_sheet = workbook["Items"]
        else:
            item_sheet = workbook.create_sheet("Items")
            item_sheet.append(ITEM_HEADERS)

    else:

        workbook = Workbook()

        invoice_sheet = workbook.active
        invoice_sheet.title = "Invoices"
        invoice_sheet.append(INVOICE_HEADERS)

        item_sheet = workbook.create_sheet("Items")
        item_sheet.append(ITEM_HEADERS)

    # ========== Getting all Invoices ==========
    invoices = data.get("invoices", [])

    # ========== Process All Invoices ==========
    for invoice_data in invoices:

        invoice = invoice_data.get("invoice", {})
        items = invoice_data.get("items", [])

        invoice_number = invoice.get("invoice_number")

        # ========== Add Invoices to Invoice Sheet ==========
        invoice_sheet.append([
            invoice_number,
            invoice.get("date"),
            invoice.get("seller"),
            invoice.get("seller_address"),
            invoice.get("seller_tax_id"),
            invoice.get("client"),
            invoice.get("client_address"),
            invoice.get("client_tax_id"),
            invoice.get("net_total"),
            invoice.get("vat"),
            invoice.get("gross_total")
        ])

        # ========== Add Items to Item Sheet ==========
        for item_number, item in enumerate(items, start=1):

            item_sheet.append([
                invoice_number,
                item_number,
                item.get("description"),
                item.get("quantity"),
                item.get("unit"),
                item.get("net_price"),
                item.get("net_worth"),
                item.get("vat_percent"),
                item.get("gross_worth")
            ])

    # ========== Save Excel Sheet ==========
    workbook.save(file_path)

    return file_path


# print(data)

output_file = save_to_excel(data, "invoices.xlsx")

# print("Excel saved:", output_file)